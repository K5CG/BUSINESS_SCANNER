"""Rigenera .docx da file .md (RIPRESA, LICENZE, …)."""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

ROOT = Path(__file__).resolve().parent.parent


def add_code_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def md_to_docx(md_path: Path, docx_path: Path) -> None:
    doc = Document()
    in_code = False
    code_lines: list[str] = []

    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                add_code_paragraph(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if line == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.LINE)
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:].strip())
            p.style = "Intense Quote"
        elif line.strip().startswith("|"):
            doc.add_paragraph(line.strip())
        elif line.strip().startswith("- [x]") or line.strip().startswith("- [X]"):
            doc.add_paragraph("✓ " + line.strip()[5:].strip(), style="List Bullet")
        elif line.strip().startswith("- [ ]"):
            doc.add_paragraph("☐ " + line.strip()[5:].strip(), style="List Bullet")
        elif line.strip().startswith("- "):
            doc.add_paragraph(line.strip()[2:].strip(), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)

    doc.save(docx_path)
    print(f"Scritto: {docx_path}")


def main() -> None:
    default_names = [
        "README",
        "RIPRESA",
        "LICENZE",
        "TEST-PLAN",
        "COSTI-E-PUBBLICAZIONE",
    ]
    names = sys.argv[1:] if len(sys.argv) > 1 else default_names
    for name in names:
        md_path = ROOT / f"{name}.md"
        docx_path = ROOT / f"{name}.docx"
        if not md_path.exists():
            print(f"Skip (manca): {md_path}")
            continue
        md_to_docx(md_path, docx_path)

    docs_dir = ROOT / "docs"
    for md_path in sorted(docs_dir.glob("*.md")):
        docx_path = docs_dir / f"{md_path.stem}.docx"
        md_to_docx(md_path, docx_path)


if __name__ == "__main__":
    main()
